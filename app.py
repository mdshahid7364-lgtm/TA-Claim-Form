from flask import Flask, render_template, request, send_from_directory
from openpyxl import load_workbook
from copy import copy
import os
import shutil

app = Flask(__name__)


# =========================================================
# HOME PAGE
# =========================================================

@app.route("/")
def home():
    return render_template("form.html")


# =========================================================
# TA FORM SUBMIT
# =========================================================

@app.route("/submit", methods=["POST"])
def submit():

    # -----------------------------------------------------
    # Employee Details
    # -----------------------------------------------------

    employee_code = request.form.get("employee_code")
    name = request.form.get("name")
    position = request.form.get("position")
    job = request.form.get("job")
    grade = request.form.get("grade")

    # -----------------------------------------------------
    # Travel Details
    # -----------------------------------------------------

    travel_date = request.form.get("travel_date")
    location_from = request.form.get("location_from")
    location_to = request.form.get("location_to")
    from_time = request.form.get("from_time")
    to_time = request.form.get("to_time")

    # -----------------------------------------------------
    # Advance Amount
    # -----------------------------------------------------

    advance = request.form.get("advance")

    if advance:
        advance = float(advance)
    else:
        advance = 0


    # -----------------------------------------------------
    # Expense Details
    # -----------------------------------------------------

    expense_dates = request.form.getlist("expense_date[]")
    particulars = request.form.getlist("particulars[]")
    modes = request.form.getlist("mode[]")
    supporting = request.form.getlist("supporting[]")
    actual_amounts = request.form.getlist("actual_amount[]")
    trs_policies = request.form.getlist("trs_policy[]")
    remarks = request.form.getlist("remarks[]")

    total_expenses = len(expense_dates)


    # =====================================================
    # EXCEL FILE SETUP
    # =====================================================

    template_file = "TA_Template.xlsx"

    # generated folder create karega
    os.makedirs("generated", exist_ok=True)

    # File name safe banane ke liye
    safe_name = name.replace(" ", "_")

    output_file = f"generated/TA_{safe_name}_{employee_code}.xlsx"

    # Original template ki copy banayega
    shutil.copy(template_file, output_file)


    # =====================================================
    # OPEN EXCEL
    # =====================================================

    workbook = load_workbook(output_file)
    sheet = workbook.active


    # =====================================================
    # EMPLOYEE DETAILS EXCEL ME FILL
    # =====================================================

    sheet["B4"] = employee_code
    sheet["D4"] = travel_date
    sheet["G4"] = ""

    sheet["B5"] = name
    sheet["D5"] = location_from
    sheet["G5"] = from_time

    sheet["B6"] = position
    sheet["D6"] = location_to
    sheet["G6"] = to_time

    sheet["B7"] = job
    sheet["D7"] = grade


    # =====================================================
    # EXPENSE ROW SETTINGS
    # =====================================================

    # Template me expense rows 9 se 22 tak hain
    expense_start_row = 9

    # Original template me 14 expense rows hain
    existing_expense_rows = 14

    # Total row normally 23 hai
    total_row_original = 23


    # =====================================================
    # AGAR EXPENSE 14 SE ZYADA HAI
    # TO EXTRA ROWS INSERT KARENGE
    # =====================================================

    if total_expenses > existing_expense_rows:

        # Kitni extra rows chahiye
        extra_rows = total_expenses - existing_expense_rows


        # -------------------------------------------------
        # Existing merged cells ko temporarily save karna
        # -------------------------------------------------

        merged_ranges_to_shift = []

        for merged_range in list(sheet.merged_cells.ranges):

            # Agar merged range total section ke neeche hai
            if merged_range.min_row >= total_row_original:

                merged_ranges_to_shift.append(
                    (
                        merged_range.min_row,
                        merged_range.max_row,
                        merged_range.min_col,
                        merged_range.max_col
                    )
                )


        # -------------------------------------------------
        # Purane merged cells ko unmerge karna
        # -------------------------------------------------

        for merged_range in list(sheet.merged_cells.ranges):

            if merged_range.min_row >= total_row_original:
                sheet.unmerge_cells(str(merged_range))


        # -------------------------------------------------
        # Extra rows insert karna
        # -------------------------------------------------

        sheet.insert_rows(
            total_row_original,
            amount=extra_rows
        )


        # -------------------------------------------------
        # New expense rows ka format copy karna
        # Row 22 = original last expense row
        # -------------------------------------------------

        source_row = 22

        for i in range(extra_rows):

            target_row = total_row_original + i

            # Row height copy
            if source_row in sheet.row_dimensions:
                sheet.row_dimensions[target_row].height = (
                    sheet.row_dimensions[source_row].height
                )

            # A se G tak formatting copy
            for col in range(1, 8):

                source_cell = sheet.cell(
                    row=source_row,
                    column=col
                )

                target_cell = sheet.cell(
                    row=target_row,
                    column=col
                )

                # Cell style copy
                target_cell._style = copy(
                    source_cell._style
                )

                # Number format
                target_cell.number_format = (
                    source_cell.number_format
                )

                # Alignment
                target_cell.alignment = copy(
                    source_cell.alignment
                )

                # Border
                target_cell.border = copy(
                    source_cell.border
                )

                # Fill
                target_cell.fill = copy(
                    source_cell.fill
                )

                # Font
                target_cell.font = copy(
                    source_cell.font
                )

                # Protection
                target_cell.protection = copy(
                    source_cell.protection
                )


        # -------------------------------------------------
        # Merged cells ko new position par restore karna
        # -------------------------------------------------

        for min_row, max_row, min_col, max_col in merged_ranges_to_shift:

            new_min_row = min_row + extra_rows
            new_max_row = max_row + extra_rows

            sheet.merge_cells(
                start_row=new_min_row,
                start_column=min_col,
                end_row=new_max_row,
                end_column=max_col
            )


    # =====================================================
    # EXPENSE DATA EXCEL ME FILL
    # =====================================================

    for i in range(total_expenses):

        row = expense_start_row + i

        # Date
        sheet[f"A{row}"] = expense_dates[i]

        # Particulars
        sheet[f"B{row}"] = particulars[i]

        # Mode
        sheet[f"C{row}"] = modes[i]

        # Supporting Documents
        sheet[f"D{row}"] = supporting[i]

        # Actual Expense
        if actual_amounts[i]:
            sheet[f"E{row}"] = float(actual_amounts[i])
        else:
            sheet[f"E{row}"] = 0

        # TRS Policy
        if i < len(trs_policies) and trs_policies[i]:
            sheet[f"F{row}"] = float(trs_policies[i])
        else:
            sheet[f"F{row}"] = 0

        # Remarks
        if i < len(remarks):
            sheet[f"G{row}"] = remarks[i]
        else:
            sheet[f"G{row}"] = ""


    # =====================================================
    # TOTAL / ADVANCE / BALANCE ROWS
    # =====================================================

    # Total row expense ke turant baad
    total_row = expense_start_row + total_expenses

    # Total Expenses row
    total_expenses_row = total_row + 1

    # Advance row
    advance_row = total_row + 2

    # Balance row
    balance_row = total_row + 3


    # -----------------------------------------------------
    # TOTAL
    # -----------------------------------------------------

    sheet[f"E{total_row}"] = (
        f"=SUM(E{expense_start_row}:E{total_row - 1})"
    )


    # -----------------------------------------------------
    # TOTAL EXPENSES
    # -----------------------------------------------------

    sheet[f"E{total_expenses_row}"] = (
        f"=E{total_row}"
    )


    # -----------------------------------------------------
    # ADVANCE
    # -----------------------------------------------------

    sheet[f"B{advance_row}"] = advance


    # -----------------------------------------------------
    # BALANCE
    # -----------------------------------------------------

    sheet[f"B{balance_row}"] = (
        f"=E{total_expenses_row}-B{advance_row}"
    )


    # =====================================================
    # SAVE EXCEL
    # =====================================================

    workbook.save(output_file)


    # =====================================================
    # DOWNLOAD FILE NAME
    # =====================================================

    filename = os.path.basename(output_file)


    # =====================================================
    # SUCCESS PAGE
    # =====================================================

    return f"""
    <!DOCTYPE html>

    <html>

    <head>

        <title>TA Submitted</title>

    </head>

    <body>

        <h1>TA Form Submitted Successfully!</h1>

        <p>
            <b>Employee Code:</b>
            {employee_code}
        </p>

        <p>
            <b>Name:</b>
            {name}
        </p>

        <p>
            <b>Job:</b>
            {job}
        </p>

        <p>
            <b>From:</b>
            {location_from}
        </p>

        <p>
            <b>To:</b>
            {location_to}
        </p>

        <br>

        <p>
            <b>
                TA Excel file has been created successfully.
            </b>
        </p>

        <br>

        <a href="/download/{filename}">

            <button>
                Download TA Excel
            </button>

        </a>

    </body>

    </html>
    """


# =========================================================
# DOWNLOAD GENERATED EXCEL
# =========================================================

@app.route("/download/<filename>")
def download_file(filename):

    return send_from_directory(
        "generated",
        filename,
        as_attachment=True
    )


# =========================================================
# START FLASK
# =========================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000
    )
